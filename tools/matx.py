"""강의자료(materials/<과목>/) → work/<SID>/mat/<파일id>/ 쪽별 추출 + 정리본 키별 쪽 이미지 연결.
사용: .venv/bin/python tools/matx.py [SID ...]        (인자 없으면 전 과목)
- t/<쪽>.txt  pymupdf 텍스트층 + PDF 주석(필기) 내용('[주석] …')
- i/<쪽>.jpg  쪽 이미지(폭 1100, 필기 주석 포함해 렌더)
- o/<쪽>.txt  tesseract OCR(kor+eng) — 텍스트층이 150자 미만인 쪽만(스캔·촬영본·이미지 슬라이드)
- pptx: t/<슬라이드>.txt(텍스트·발표자 노트), i/ 없음(그림은 pptx 안의 이미지만 p<슬라이드>_<n>.<ext>로 저장)
- index.json  파일별 쪽수·원본 파일명
- work/<SID>/lec/<폴더> → mat/<파일id>/i 연결(subject.py LECMAP 폴더 이름 = SRC 표) — build4가 이 쪽 이미지를 씀"""
import os, sys, re, json, subprocess, unicodedata
from multiprocessing import Pool
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import jblpaths as J
import pymupdf

SUBJ = {'구강악안면외과학1': 'OMS1', '임상치과보존학': 'CONS', '치과임플란트학': 'IMPL', '임상두경부해부학': 'ANAT',
        '노인치과학': 'GERI', '임상치과약물치료학': 'PHARM', '심미치과학': 'ESTH'}
# 정리본 키의 쪽 이미지 폴더(subject.LECMAP[키][0], OMS1은 끝에 'i'가 붙음) ← 강의자료 파일 앞머리
SRC = {
    'OMS1': {'L08i': '260831_', 'L09i': '260907_', 'L10i': '260914_', 'L03i': '0922_', 'L04i': '0929_', 'L05i': '1013_', 'L06i': '1020(1)_', 'L07i': '1020(2)_'},
    'CONS': {'C09i': '260908_', 'C02i': '250909(1)_', 'C03i': '250909(2)_', 'C04i': '250923_', 'C05i': '250930_', 'C06i': '251014_', 'C07i': '251021_', 'C08i': '260901_'},
    'IMPL': {'I01i': '260901_', 'I02i': '250909_', 'I03i': '260915_', 'I04i': '250916_명훈_ONJ', 'I05i': '250923_', 'I06i': '260908_', 'I07i': '250930_', 'I08i': '251014_', 'I09i': '251021_', 'I10i': '250902_'},
    'ANAT': {'A01i': '260902_', 'A02i': '260909_', 'A03i': '250915_', 'A04i': '260916_', 'A05i': '250917_', 'A06i': '250924_', 'A07i': '251001_', 'A08i': '251022_', 'A09i': '251027_', 'A10i': '250903_'},
    'GERI': {'G01i': '250904_', 'G02i': '250911_', 'G03i': '260903_', 'G04i': '250925_', 'G05i': '251002_', 'G06i': '260917_', 'G07i': '251016_', 'G08i': '251023_'},
    'PHARM': {'P01i': '260903_', 'P02i': '250904_', 'P03i': '260910_', 'P04i': '250911_', 'P05i': '260917_', 'P06i': '250925_', 'P07i': '251002_', 'P08i': '251016_', 'P09i': '251023_', 'P10i': '250918_'},
    # ESTH: 강의 주제마다 가장 최신 연도(26 → 없으면 25). E08i~E10i는 26 강의의 25년도 판(보조 키 — 25 필기·그림)
    'ESTH': {'E01i': '260911_이창하_Introduction', 'E02i': '260911_이창하_Fundamentals', 'E03i': '260918_', 'E04i': '250926_', 'E05i': '251010_',
             'E06i': '251017_', 'E07i': '251024_', 'E08i': '250905_', 'E09i': '250912_', 'E10i': '250919_'},
}
nfc = lambda s: unicodedata.normalize('NFC', s)
def fid(fn): return re.sub(r'[^0-9A-Za-z가-힣()_.-]+', '_', nfc(os.path.splitext(fn)[0]))[:80]

def page_job(a):
    src, out, i, need_ocr = a
    doc = pymupdf.open(src); pg = doc[i - 1]
    t = pg.get_text()
    notes = []
    for an in (pg.annots() or []):
        c = (an.info.get('content') or '').strip()
        if c: notes.append('[주석] ' + c.replace('\n', ' '))
    open(f'{out}/t/{i}.txt', 'w', encoding='utf-8').write(t + ('\n' + '\n'.join(notes) if notes else ''))
    if not os.path.exists(f'{out}/i/{i}.jpg'):
        z = 1100 / pg.rect.width
        pg.get_pixmap(matrix=pymupdf.Matrix(z, z), annots=True).pil_save(f'{out}/i/{i}.jpg', quality=82)
    if need_ocr and not os.path.exists(f'{out}/o/{i}.txt'):
        z = min(2200 / pg.rect.width, 4.0); tmp = f'{out}/o/_{i}.png'
        pg.get_pixmap(matrix=pymupdf.Matrix(z, z), annots=True).save(tmp)
        r = subprocess.run(['tesseract', tmp, 'stdout', '-l', 'kor+eng', '--psm', '3'], capture_output=True)
        open(f'{out}/o/{i}.txt', 'w', encoding='utf-8').write(r.stdout.decode('utf-8', 'replace')); os.remove(tmp)
    return i

def pptx_job(src, out):
    from pptx import Presentation
    pr = Presentation(src); n = 0
    for i, sl in enumerate(pr.slides, 1):
        n = i; lines, k = [], 0
        for sh in sl.shapes:
            if sh.has_text_frame:
                for para in sh.text_frame.paragraphs:
                    t = ''.join(r.text for r in para.runs).strip()
                    if t: lines.append(('  ' * para.level) + t)
            if getattr(sh, 'has_table', False) and sh.has_table:
                for row in sh.table.rows: lines.append(' | '.join(c.text.strip() for c in row.cells))
            if sh.shape_type == 13:
                k += 1; im = sh.image; open(f'{out}/i/p{i}_{k}.{im.ext}', 'wb').write(im.blob)
        if sl.has_notes_slide and sl.notes_slide.notes_text_frame.text.strip():
            lines.append('[발표자 노트] ' + sl.notes_slide.notes_text_frame.text.strip())
        open(f'{out}/t/{i}.txt', 'w', encoding='utf-8').write('\n'.join(lines))
    return n

def run(sid, pool):
    sub = next(k for k, v in SUBJ.items() if v == sid); d = os.path.join(J.ROOT, 'materials', sub)
    if not os.path.isdir(d):
        d = next((os.path.join(J.ROOT, 'materials', x) for x in os.listdir(os.path.join(J.ROOT, 'materials')) if nfc(x) == sub), None)
    if not d: return
    idx = {}; mat = J.work(sid, 'mat')
    for fn in sorted(os.listdir(d)):
        low = fn.lower(); src = os.path.join(d, fn)
        if not (low.endswith('.pdf') or low.endswith('.pptx')): continue
        out = os.path.join(mat, fid(fn))
        for s in ('t', 'i', 'o'): os.makedirs(f'{out}/{s}', exist_ok=True)
        if low.endswith('.pptx'):
            n = pptx_job(src, out); idx[fid(fn)] = {'file': nfc(fn), 'pages': n, 'kind': 'pptx'}; continue
        doc = pymupdf.open(src); n = len(doc)
        jobs = [(src, out, i, len(doc[i - 1].get_text().strip()) < 150) for i in range(1, n + 1)]
        pool.map(page_job, jobs, chunksize=2)
        idx[fid(fn)] = {'file': nfc(fn), 'pages': n, 'kind': 'pdf', 'ocr': sum(1 for j in jobs if j[3])}
        print(f'  {sid} {n:4}쪽 OCR {idx[fid(fn)]["ocr"]:3}  {nfc(fn)}', flush=True)
    json.dump(idx, open(os.path.join(mat, 'index.json'), 'w'), ensure_ascii=False, indent=1)
    # 정리본 키 폴더 연결
    lec = J.work(sid, 'lec'); os.makedirs(lec, exist_ok=True)
    for folder, pre in SRC.get(sid, {}).items():
        hit = [k for k, v in idx.items() if v['file'].startswith(pre)]
        if len(hit) != 1: print(f'  ✗ {sid} {folder}: {pre} → {hit}'); continue
        link = os.path.join(lec, folder)
        if os.path.islink(link) or os.path.exists(link):
            if os.path.islink(link): os.remove(link)
            else: continue
        os.symlink(os.path.join('..', 'mat', hit[0], 'i'), link)

if __name__ == '__main__':
    want = [a.upper() for a in sys.argv[1:]] or list(SRC)
    with Pool(7) as pool:
        for sid in want: run(sid, pool)
    print('DONE')
